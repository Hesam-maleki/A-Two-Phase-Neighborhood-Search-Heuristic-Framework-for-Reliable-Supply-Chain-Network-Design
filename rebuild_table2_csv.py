"""
rebuild_table2_csv.py: rebuild table2_combined.csv for an existing run folder
from its saved COMPARE_*.xlsx workbooks, without re-running any experiment.

Usage:
    python rebuild_table2_csv.py run_<timestamp>
"""
import os
import sys
import glob

from openpyxl import load_workbook

from reproduce_paper_table2 import (
    INSTANCES,
    RESULTS_ROOT,
    COMBINED_CSV_NAME,
    _row_from_results,
    _write_combined_csv,
)

CASES = ["case1", "case2", "case3", "case4"]


def _summary_rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = {r[0]: r[1:5] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    wb.close()
    return rows


def _results_from_workbook(path):
    s = _summary_rows(path)
    results = {}
    for i, case in enumerate(CASES):
        results[case] = dict(
            status    = s["Status"][i],
            time      = float(s["Solve time (s)"][i]),
            iters     = s["Iterations"][i],
            p2_sweeps = s["Phase 2 sweeps (C3/C4)"][i],
            UB        = float(s["Upper bound (UB)"][i]),
            LB        = float(s["Lower bound (LB)"][i]),
            gap       = float(s["Gap %"][i]),
        )
    return results


def _workbook_for(run_dir, label):
    safe_label = label.replace("%", "pct")
    matches = glob.glob(os.path.join(run_dir, f"COMPARE_{safe_label}_seed*_m*_*.xlsx"))
    if len(matches) != 1:
        sys.exit(f"expected one workbook for {label} in {run_dir}, found {len(matches)}")
    return matches[0]


def main():
    if len(sys.argv) != 2:
        sys.exit("usage: python rebuild_table2_csv.py run_<timestamp>")
    run_dir = sys.argv[1]
    if not os.path.isdir(run_dir):
        run_dir = os.path.join(RESULTS_ROOT, os.path.basename(run_dir.rstrip("/\\")))
    if not os.path.isdir(run_dir):
        sys.exit(f"no such run folder: {run_dir}")

    rows_by_label = {}
    for label, *_ in INSTANCES:
        results = _results_from_workbook(_workbook_for(run_dir, label))
        rows_by_label[label] = _row_from_results(label, results)

    combined_csv = os.path.join(run_dir, COMBINED_CSV_NAME)
    _write_combined_csv(rows_by_label, combined_csv)
    print(f"Rebuilt {combined_csv} from {len(rows_by_label)} workbooks")


if __name__ == "__main__":
    main()
