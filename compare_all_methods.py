"""
compare_all_methods.py: run the four methods of Table 2 on one instance and
export a comparison workbook.

Case 1: Exact C&CG
Case 2: Early-stopped C&CG, taken from Phase 1 of the H1 run (not a separate run)
Case 3: H1, early-stopped C&CG + Phase 2
Case 4: H2, no-disruption MIP + Phase 2
"""

import math
import time
import datetime

from ro_lndp0_ccg import (
    LNDPInstance,
    compute_big_M,
    solve_subproblem,
    build_master,
    add_ccg_cut,
    solve_mp_lp,
)
from h1_ccg_ns import h1_ccg_ns
from h2_mip_ns import h2_mip_ns
from instance_generator_paper_exact import (
    generate_instance_paper_exact,
    summarize_instance,
)


# Recourse helpers

def compute_worst_case(inst: LNDPInstance, y: dict,
                       big_M: dict, time_limit: float = 300.0):
    """Return (z_hat, recourse_cost, fixed_cost, total_cost) for y."""
    z_hat, chi = solve_subproblem(inst, y, big_M, time_limit)
    fc = sum(inst.f[j] * y.get(j, 0) for j in inst.V0)
    return z_hat, chi, fc, fc + chi


def compute_recourse_for_scenario(inst: LNDPInstance, y: dict, z: dict):
    """
    Solve the recourse LP for fixed y and disruption z.
    Returns (recourse_cost, unmet_total, flow_vals, unmet_vals).
    """
    from docplex.mp.model import Model
    lp = Model(name="recourse", log_output=False)

    x = {(i, j): lp.continuous_var(lb=0) for (i, j) in inst.arcs}
    u = {j: lp.continuous_var(lb=0) for j in inst.VD}

    for j in inst.VS:
        lp.add_constraint(lp.sum(x[(j, k)] for k in inst.Vplus[j]) <= inst.b[j])
    for j in inst.VT:
        lp.add_constraint(
            lp.sum(x[(i, j)] for i in inst.Vminus[j])
            == lp.sum(x[(j, k)] for k in inst.Vplus[j]))
    for j in inst.VD:
        lp.add_constraint(
            lp.sum(x[(i, j)] for i in inst.Vminus[j]) + u[j] == -inst.b[j])
    for j in inst.V0:
        cap = (1 - z.get(j, 0)) * inst.Q[j] * y.get(j, 0)
        lp.add_constraint(lp.sum(x[(j, k)] for k in inst.Vplus[j]) <= cap)

    lp.minimize(
        lp.sum(inst.c[(i, j)] * x[(i, j)] for (i, j) in inst.arcs)
        + lp.sum(inst.theta[j] * u[j] for j in inst.VD))

    sol = lp.solve(log_output=False)
    if not sol:
        lp.end()
        return math.inf, math.inf, {}, {}

    recourse = lp.objective_value
    unmet_total = sum(u[j].solution_value for j in inst.VD)
    flow_vals = {arc: x[arc].solution_value
                 for arc in inst.arcs if x[arc].solution_value > 1e-6}
    unmet_vals = {j: u[j].solution_value
                  for j in inst.VD if u[j].solution_value > 1e-6}
    lp.end()
    return recourse, unmet_total, flow_vals, unmet_vals


def build_flow_tables(inst: LNDPInstance, flow_vals: dict):
    """
    Build two structured tables from a raw arc-flow solution.

    flow_to_demand : list of dicts, one row per (origin -> demand_node) arc
    origin_summary : list of dicts, one row per origin sending flow to demand
    """
    from collections import defaultdict
    VD_set = set(inst.VD)
    VS_set = set(inst.VS)
    VT_set = set(inst.VT)

    demand_arcs = {(i, j): v for (i, j), v in flow_vals.items() if j in VD_set}

    flow_to_demand = []
    for (i, j), flow in sorted(demand_arcs.items()):
        unit_cost = inst.c.get((i, j), 0.0)
        flow_to_demand.append(dict(
            origin      = i,
            dest        = j,
            flow        = round(flow, 4),
            unit_cost   = round(unit_cost, 4),
            arc_cost    = round(flow * unit_cost, 4),
            origin_type = ("supply"    if i in VS_set else
                           "transship" if i in VT_set else "other"),
        ))

    by_origin = defaultdict(lambda: {"total_flow": 0.0,
                                     "total_cost": 0.0, "serves": {}})
    for row in flow_to_demand:
        o = row["origin"]
        by_origin[o]["total_flow"] += row["flow"]
        by_origin[o]["total_cost"] += row["arc_cost"]
        by_origin[o]["serves"][row["dest"]] = row["flow"]

    origin_summary = []
    for node in sorted(by_origin.keys()):
        d = by_origin[node]
        origin_summary.append(dict(
            node       = node,
            node_type  = "supply" if node in VS_set else "transship",
            total_flow = round(d["total_flow"], 4),
            total_cost = round(d["total_cost"], 4),
            serves     = dict(d["serves"]),
        ))

    return flow_to_demand, origin_summary


# Cases 1 & 2: C&CG with configurable stop gap

def run_ccg_early_stop(inst: LNDPInstance,
                       big_M: dict,
                       stop_gap_pct: float = 5.0,
                       time_limit: float = 10800.0,
                       eps: float = 1e-4,
                       use_lp_enhancement: bool = True,
                       verbose: bool = True,
                       label: str = ""):
    """
    Standard C&CG stopped at stop_gap_pct (or at full convergence if 0).
    stop_gap_pct = 0.0 => exact C&CG (Case 1).
    stop_gap_pct > 0   => early-stop C&CG (Case 2).
    """
    t_start = time.time()
    LB = -math.inf
    UB = math.inf
    y_star = None
    history = []
    r = 0
    status = "EARLY_STOP"

    mp, y, phi = build_master(inst)
    y_hat = {j: 1 for j in inst.V0}

    if verbose:
        print("\n" + "=" * 72)
        tag = f"  [{label}]" if label else ""
        print(f"  C&CG  stop_gap={stop_gap_pct:.1f}%{tag}")
        print("=" * 72)
        hdr = (f"{'Iter':>4}  {'LB':>14}  {'UB':>14}  {'Gap%':>8}  "
               f"{'SubP(s)':>8}  {'MP(s)':>8}")
        print(hdr)
        print("-" * len(hdr))

    while True:
        elapsed = time.time() - t_start
        if elapsed >= time_limit:
            status = "TIME_LIMIT"
            break

        iter_t0 = time.time()
        remaining = max(10.0, time_limit - elapsed)

        subp_t0 = time.time()
        z_hat, chi = solve_subproblem(inst, y_hat, big_M, remaining)
        subp_time = time.time() - subp_t0
        r += 1

        add_ccg_cut(mp, inst, y, phi, z_hat, r)

        fc_hat = sum(inst.f[j] * y_hat[j] for j in inst.V0)
        UB_cand = fc_hat + chi
        if UB_cand < UB:
            UB = UB_cand
            y_star = dict(y_hat)

        master_t0 = time.time()
        if use_lp_enhancement:
            lb_lp = solve_mp_lp(mp, y)
            LB = max(LB, lb_lp)
        sol_mp = mp.solve(log_output=False)
        master_time = time.time() - master_t0

        if sol_mp:
            LB = max(LB, mp.objective_value)
            y_hat = {j: int(round(y[j].solution_value)) for j in inst.V0}

        gap = (UB - LB) / abs(UB) * 100 if UB not in (0.0, math.inf) else math.inf
        elapsed = time.time() - t_start

        history.append(dict(
            iter         = r,
            LB           = LB,
            UB           = UB,
            gap          = gap,
            time_elapsed = elapsed,
            subp_time    = subp_time,
            master_time  = master_time,
            iter_time    = time.time() - iter_t0,
            z_hat        = dict(z_hat),
            y_hat        = dict(y_hat),
        ))

        if verbose:
            print(f"{r:>4}  {LB:>14.2f}  {UB:>14.2f}  {gap:>7.4f}%  "
                  f"{subp_time:>8.2f}  {master_time:>8.2f}")

        if gap <= eps * 100:
            status = "OPTIMAL"
            if verbose:
                print(f"\n  Converged to optimality (gap={gap:.6f}%)")
            break

        if stop_gap_pct > 0 and gap <= stop_gap_pct:
            status = "EARLY_STOP"
            if verbose:
                print(f"\n  Early stop: gap {gap:.4f}% <= {stop_gap_pct}%")
            break

    mp.end()
    return y_star, UB, LB, history, status


# Comparison runner

def run_comparison(inst: LNDPInstance,
                   raw: dict,
                   gen_params: dict,
                   stop_gap_pct: float = 10.0,
                   time_limit_exact: float = 10800.0,
                   time_limit_H1: float = 10800.0,
                   time_limit_H2: float = 10800.0,
                   eps_exact: float = 1e-6,
                   pool_cap_H1: int = 0,
                   pool_cap_H2: int = 0,
                   verbose: bool = True):
    """Run the four cases on one instance and return a results dict."""
    big_M = compute_big_M(inst)
    results = {}

    # Case 1: Exact C&CG
    print("\n" + "#" * 72)
    print("  CASE 1 - Exact C&CG  (full convergence)")
    print("#" * 72)
    t0 = time.time()
    y1, UB1, LB1, hist1, st1 = run_ccg_early_stop(
        inst, big_M,
        stop_gap_pct       = 0.0,
        time_limit         = time_limit_exact,
        eps                = eps_exact,
        use_lp_enhancement = True,
        verbose            = verbose,
        label              = "Case 1 Exact",
    )
    time1 = time.time() - t0
    z1, chi1, fc1, _ = compute_worst_case(inst, y1, big_M)
    rec1, unmet1, flow1, unmet_vals1 = compute_recourse_for_scenario(inst, y1, z1)
    ftd1, ors1 = build_flow_tables(inst, flow1)

    results["case1"] = dict(
        label          = "Exact C&CG",
        y              = y1,
        UB             = UB1,
        LB             = LB1,
        gap            = (UB1 - LB1) / abs(UB1) * 100 if UB1 > 0 else math.inf,
        fc             = fc1,
        recourse       = chi1,
        time           = time1,
        iters          = len(hist1),
        status         = st1,
        z_worst        = z1,
        unmet          = unmet1,
        unmet_vals     = unmet_vals1,
        flow_vals      = flow1,
        flow_to_demand = ftd1,
        origin_summary = ors1,
        history        = hist1,
    )

    # Case 3: H1
    print("\n" + "#" * 72)
    print(f"  CASE 3 - H1  (Phase 1 early-stop at {stop_gap_pct}%)")
    print("#" * 72)
    t0 = time.time()
    y3, UB3, LB3, hist3, st3, phase1_3 = h1_ccg_ns(
        inst,
        switch_gap_pct     = stop_gap_pct,
        phase1_time_limit  = time_limit_H1,
        phase2_time_limit  = time_limit_H1,
        eps                = eps_exact,
        use_lp_enhancement = True,
        pool_cap           = pool_cap_H1,
        verbose            = verbose,
    )
    time3 = time.time() - t0
    z3, chi3, fc3, _ = compute_worst_case(inst, y3, big_M)
    rec3, unmet3, flow3, unmet_vals3 = compute_recourse_for_scenario(inst, y3, z3)
    ftd3, ors3 = build_flow_tables(inst, flow3)

    p1_iters3 = sum(1 for h in hist3 if h.get("phase", 1) == 1)
    p2_sweeps3 = sum(1 for h in hist3 if h.get("phase", 1) == 2)

    results["case3"] = dict(
        label          = "H1 (C&CG + Phase 2)",
        y              = y3,
        UB             = UB3,
        LB             = LB3,
        gap            = (UB3 - LB3) / abs(UB3) * 100 if UB3 > 0 else math.inf,
        fc             = fc3,
        recourse       = chi3,
        time           = time3,
        p1_iters       = p1_iters3,
        p2_sweeps      = p2_sweeps3,
        iters          = len(hist3),
        status         = st3,
        z_worst        = z3,
        unmet          = unmet3,
        unmet_vals     = unmet_vals3,
        flow_vals      = flow3,
        flow_to_demand = ftd3,
        origin_summary = ors3,
        history        = hist3,
    )

    # Case 2: the Phase 1 result of the H1 run above, not a separate run
    y2  = dict(phase1_3["y"])
    UB2 = phase1_3["UB"]
    LB2 = phase1_3["LB"]
    time2 = phase1_3["time"]
    st2 = phase1_3["status"]
    z2, chi2, fc2, _ = compute_worst_case(inst, y2, big_M)
    rec2, unmet2, flow2, unmet_vals2 = compute_recourse_for_scenario(inst, y2, z2)
    ftd2, ors2 = build_flow_tables(inst, flow2)
    results["case2"] = dict(
        label          = "Early-stopped C&CG (Phase 1 of H1)",
        y              = y2,
        UB             = UB2,
        LB             = LB2,
        gap            = (UB2 - LB2) / abs(UB2) * 100 if UB2 > 0 else math.inf,
        fc             = fc2,
        recourse       = chi2,
        time           = time2,
        iters          = phase1_3["iters"],
        status         = st2,
        z_worst        = z2,
        unmet          = unmet2,
        unmet_vals     = unmet_vals2,
        flow_vals      = flow2,
        flow_to_demand = ftd2,
        origin_summary = ors2,
        history        = phase1_3["history"],
    )

    # Case 4: H2
    print("\n" + "#" * 72)
    print("  CASE 4 - H2  (no-disruption MIP + Phase 2)")
    print("#" * 72)
    t0 = time.time()
    y4, UB4, LB4, hist_p1_4, hist_p2_4, st4 = h2_mip_ns(
        inst,
        phase2_time_limit = time_limit_H2,
        pool_cap          = pool_cap_H2,
        verbose           = verbose,
    )
    time4 = time.time() - t0
    z4, chi4, fc4, _ = compute_worst_case(inst, y4, big_M)
    rec4, unmet4, flow4, unmet_vals4 = compute_recourse_for_scenario(inst, y4, z4)
    ftd4, ors4 = build_flow_tables(inst, flow4)

    case4_history = [dict(
        iter         = 0,
        LB           = LB4,
        UB           = hist_p1_4.get("UB_start", UB4),
        gap          = ((hist_p1_4.get("UB_start", UB4) - LB4)
                        / abs(hist_p1_4.get("UB_start", UB4)) * 100
                        if hist_p1_4.get("UB_start", UB4) > 0 else math.inf),
        time_elapsed = hist_p1_4.get("mip_time", 0) + hist_p1_4.get("subp_time", 0),
        subp_time    = hist_p1_4.get("subp_time", 0),
        master_time  = hist_p1_4.get("mip_time", 0),
        iter_time    = hist_p1_4.get("mip_time", 0) + hist_p1_4.get("subp_time", 0),
        z_hat        = dict(hist_p1_4.get("z_init", {})),
        y_hat        = dict(hist_p1_4.get("y_start", {})),
        phase        = 1,
        move_type    = "Phase 1 (no-disruption MIP)",
    )]
    case4_history.extend(hist_p2_4)

    results["case4"] = dict(
        label          = "H2 (MIP + Phase 2)",
        y              = y4,
        UB             = UB4,
        LB             = LB4,
        gap            = (UB4 - LB4) / abs(UB4) * 100 if UB4 > 0 else math.inf,
        fc             = fc4,
        recourse       = chi4,
        time           = time4,
        p1_iters       = 1,
        p2_sweeps      = len(hist_p2_4),
        iters          = 1 + len(hist_p2_4),
        status         = st4,
        z_worst        = z4,
        unmet          = unmet4,
        unmet_vals     = unmet_vals4,
        flow_vals      = flow4,
        flow_to_demand = ftd4,
        origin_summary = ors4,
        history        = case4_history,
    )

    return results


# Console comparison report

def print_comparison(results: dict, inst: LNDPInstance):
    c1, c2, c3, c4 = (results["case1"], results["case2"],
                     results["case3"], results["case4"])
    cases = [c1, c2, c3, c4]
    W = 22

    def row(label, vals, fmt="{}", suffix=""):
        line = f"  {label:<32}"
        for v in vals:
            cell = (fmt.format(v) + suffix) if v is not None else "-"
            line += f"{cell:>{W}}"
        print(line)

    def sep():
        print("  " + "-" * (32 + W * 4))

    labels = ["Exact C&CG", "Early-stopped C&CG", "H1", "H2"]
    print("\n" + "=" * (32 + W * 4 + 2))
    print("  FOUR-METHOD COMPARISON")
    print("=" * (32 + W * 4 + 2))
    line = f"  {'':32}"
    for lbl in labels:
        line += f"{lbl:>{W}}"
    print(line)
    sep()

    row("Status",             [c["status"] for c in cases])
    row("Iterations / sweeps",[c1["iters"], c2["iters"],
                               c3["p2_sweeps"], c4["p2_sweeps"]])
    row("Solve time (s)",     [c["time"] for c in cases], fmt="{:.1f}")
    sep()
    row("Upper bound (UB)",   [c["UB"] for c in cases],   fmt="{:,.2f}")
    row("Lower bound (LB)",   [c["LB"] for c in cases],   fmt="{:,.2f}")
    row("Gap %",              [max(0.0, c["gap"]) for c in cases], fmt="{:.4f}", suffix="%")
    sep()
    row("Fixed cost",         [c["fc"] for c in cases],   fmt="{:,.2f}")
    row("Worst-case recourse",[c["recourse"] for c in cases], fmt="{:,.2f}")
    sep()
    row("Unmet demand",       [c["unmet"] for c in cases], fmt="{:.2f}")

    print()
    for lbl, c in zip(labels, cases):
        open_facs = sorted(j for j in inst.V0 if c["y"].get(j, 0) == 1)
        print(f"  {lbl:<32}  open: {open_facs}")

    print("=" * (32 + W * 4 + 2))


# Excel export

def export_comparison_excel(results: dict, inst: LNDPInstance,
                            raw: dict, gen_params: dict,
                            out_path: str):
    """Write a 9-sheet Excel workbook comparing all four methods."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.chart import LineChart, Reference
    except ImportError:
        print("  [openpyxl not installed, skipping Excel export]")
        return

    c1, c2, c3, c4 = (results["case1"], results["case2"],
                     results["case3"], results["case4"])
    cases = [c1, c2, c3, c4]

    HDR_FILL  = PatternFill("solid", fgColor="1F4E79")
    C1_FILL   = PatternFill("solid", fgColor="DEEAF1")
    C2_FILL   = PatternFill("solid", fgColor="E2EFDA")
    C3_FILL   = PatternFill("solid", fgColor="FFF2CC")
    C4_FILL   = PatternFill("solid", fgColor="FCE4D6")
    ALT_FILL  = PatternFill("solid", fgColor="F5F5F5")
    GOOD_FILL = PatternFill("solid", fgColor="C6EFCE")
    BAD_FILL  = PatternFill("solid", fgColor="FFCCCC")
    case_fills = [C1_FILL, C2_FILL, C3_FILL, C4_FILL]

    HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
    BOLD     = Font(bold=True, size=10)
    NORM     = Font(size=10)
    thin     = Side(style="thin", color="CCCCCC")
    BORDER   = Border(left=thin, right=thin, top=thin, bottom=thin)
    center   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left     = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    NUM2     = '#,##0.00'
    NUM4     = '#,##0.0000'
    NUM6     = '0.000000'

    def _set(ws, r, c, val, font=None, fill=None, align=None, num_fmt=None):
        cl = ws.cell(row=r, column=c, value=val)
        cl.border = BORDER
        if font:    cl.font          = font
        if fill:    cl.fill          = fill
        if align:   cl.alignment     = align
        if num_fmt: cl.number_format = num_fmt
        return cl

    def _hrow(ws, r, labels, fills=None):
        for ci, lbl in enumerate(labels, 1):
            f = fills[ci - 1] if fills and ci - 1 < len(fills) else HDR_FILL
            _set(ws, r, ci, lbl, font=HDR_FONT, fill=f, align=center)
        ws.row_dimensions[r].height = 22

    def _auto_col(ws):
        for col in ws.columns:
            mx = max((len(str(cl.value or "")) for cl in col), default=8)
            ws.column_dimensions[
                get_column_letter(col[0].column)].width = min(mx + 4, 40)

    instance_label = (
        f"{int(gen_params['density']*100)}%-"
        f"{gen_params['nS']}-{gen_params['nT']}-{gen_params['nD']}  "
        f"seed={gen_params['seed']}  m={gen_params['m_budget']}"
    )

    wb = openpyxl.Workbook()

    # SHEET 1: Summary
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.freeze_panes = "B3"

    ws1.merge_cells("A1:E1")
    cl = ws1.cell(row=1, column=1,
                  value=f"Four-Method Comparison  |  {instance_label}")
    cl.font      = Font(bold=True, size=13, color="FFFFFF")
    cl.fill      = HDR_FILL
    cl.alignment = center
    ws1.row_dimensions[1].height = 28

    case_labels = [c["label"] for c in cases]
    _hrow(ws1, 2,
          ["Metric"] + case_labels,
          fills=[HDR_FILL] + case_fills)

    summary_rows = [
        ("Status",                 [c["status"] for c in cases], None),
        ("Solve time (s)",         [round(c["time"], 2) for c in cases], NUM2),
        ("Iterations",             [c.get("iters") for c in cases], None),
        ("Phase 1 iters (C3/C4)",  [None, None,
                                    c3.get("p1_iters"),
                                    c4.get("p1_iters")], None),
        ("Phase 2 sweeps (C3/C4)", [None, None,
                                    c3.get("p2_sweeps"),
                                    c4.get("p2_sweeps")], None),
        (None, None, None),
        ("Upper bound (UB)",       [round(c["UB"], 2) for c in cases], NUM2),
        ("Lower bound (LB)",       [round(c["LB"], 2) for c in cases], NUM2),
        ("Gap %",                  [round(c["gap"], 6) for c in cases], NUM6),
        ("Fixed cost",             [round(c["fc"], 2) for c in cases], NUM2),
        ("Worst-case recourse",    [round(c["recourse"], 2) for c in cases], NUM2),
        ("Unmet demand (units)",   [round(c["unmet"], 2) for c in cases], NUM2),
        (None, None, None),
        ("Open facilities",
         [str(sorted(j for j in inst.V0 if c["y"].get(j, 0) == 1))
          for c in cases], None),
        ("Disrupted facilities",
         [str(sorted(j for j in inst.V0 if c["z_worst"].get(j, 0) == 1))
          for c in cases], None),
    ]

    ri = 3
    for item in summary_rows:
        if item[0] is None:
            ri += 1
            continue
        metric, vals, fmt = item
        bg = ALT_FILL if ri % 2 == 0 else None
        _set(ws1, ri, 1, metric, font=BOLD, fill=bg, align=left)
        for ci, (v, cf) in enumerate(zip(vals, case_fills), 2):
            _set(ws1, ri, ci, v, font=NORM,
                 fill=cf if v is not None else bg,
                 align=center, num_fmt=fmt)
        ws1.row_dimensions[ri].height = 18
        ri += 1

    ws1.column_dimensions["A"].width = 30
    for ci in range(2, 6):
        ws1.column_dimensions[get_column_letter(ci)].width = 28

    # SHEETS 2-5: Convergence histories
    hist_sheets = [
        (wb.create_sheet("Case 1 - C&CG History"),  c1, C1_FILL),
        (wb.create_sheet("Case 2 - Early-stop"),    c2, C2_FILL),
        (wb.create_sheet("Case 3 - H1"), c3, C3_FILL),
        (wb.create_sheet("Case 4 - H2"), c4, C4_FILL),
    ]

    for ws_h, case, hfill in hist_sheets:
        ws_h.freeze_panes = "A2"
        cols = ["Iter", "LB", "UB", "Gap %",
                "SubP (s)", "MP (s)", "Iter (s)", "Elapsed (s)",
                "Open facilities (y_hat)",
                "Disrupted (z_hat)"]
        _hrow(ws_h, 1, cols, fills=[hfill] * len(cols))

        for ri, h in enumerate(case["history"], 2):
            z = h.get("z_hat", {})
            yh = h.get("y_hat", {})
            dis = sorted(j for j in inst.V0 if z.get(j, 0) == 1)
            open_ = sorted(j for j in inst.V0 if yh.get(j, 0) == 1)
            bg = ALT_FILL if ri % 2 == 0 else None
            ph = h.get("phase", 1)
            row_fill = (PatternFill("solid", fgColor="FFF2CC")
                        if ph == 2 else bg)
            vals = [
                h.get("iter"),
                h.get("LB"),
                h.get("UB"),
                h.get("gap"),
                round(h.get("subp_time",    0), 3),
                round(h.get("master_time",  0), 3),
                round(h.get("iter_time",    0), 3),
                round(h.get("time_elapsed", 0), 3),
                str(open_),
                str(dis),
            ]
            fmts = [None, NUM2, NUM2, NUM6, NUM2, NUM2, NUM2, NUM2, None, None]
            for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
                _set(ws_h, ri, ci, v, fill=row_fill, align=center, num_fmt=fmt)
            ws_h.row_dimensions[ri].height = 16

        try:
            chart = LineChart()
            chart.title        = f"{case['label']} - LB and UB convergence"
            chart.y_axis.title = "Cost"
            chart.x_axis.title = "Iteration"
            chart.style        = 10
            chart.width        = 20
            chart.height       = 12
            n = len(case["history"])
            lb_ref = Reference(ws_h, min_col=2, min_row=1, max_row=n + 1)
            ub_ref = Reference(ws_h, min_col=3, min_row=1, max_row=n + 1)
            chart.add_data(lb_ref, titles_from_data=True)
            chart.add_data(ub_ref, titles_from_data=True)
            chart.series[0].graphicalProperties.line.solidFill = "2E75B6"
            chart.series[1].graphicalProperties.line.solidFill = "C00000"
            ws_h.add_chart(chart, f"A{n + 5}")
        except Exception:
            pass

        _auto_col(ws_h)

    # SHEET 6: Facility decisions
    ws5 = wb.create_sheet("Facility Decisions")
    ws5.freeze_panes = "A2"

    cols5 = ["Facility", "Type",
             "Fixed cost", "Capacity",
             "Open C1", "Disrupt C1",
             "Open C2", "Disrupt C2",
             "Open C3", "Disrupt C3",
             "Open C4", "Disrupt C4"]
    _hrow(ws5, 1, cols5,
          fills=[HDR_FILL] * 4
                + [C1_FILL] * 2 + [C2_FILL] * 2
                + [C3_FILL] * 2 + [C4_FILL] * 2)

    VS_set = set(inst.VS)
    for ri, j in enumerate(sorted(inst.V0), 2):
        bg = ALT_FILL if ri % 2 == 0 else None
        ftype = "Supply" if j in VS_set else "Transshipment"
        vals = [j, ftype, round(inst.f[j], 2), round(inst.Q[j], 2)]
        fmts = [None, None, NUM2, NUM2]
        for c in cases:
            is_open = c["y"].get(j, 0) == 1
            is_dis  = c["z_worst"].get(j, 0) == 1
            vals += ["OPEN" if is_open else "closed",
                     "FAIL" if is_dis  else "ok"]
            fmts += [None, None]

        for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
            fill = bg
            if ci in (5, 7, 9, 11):
                case_idx = (ci - 5) // 2
                fill = GOOD_FILL if cases[case_idx]["y"].get(j, 0) == 1 else BAD_FILL
            elif ci in (6, 8, 10, 12):
                case_idx = (ci - 6) // 2
                fill = BAD_FILL if cases[case_idx]["z_worst"].get(j, 0) == 1 else GOOD_FILL
            _set(ws5, ri, ci, v, fill=fill, align=center, num_fmt=fmt)
        ws5.row_dimensions[ri].height = 16

    _auto_col(ws5)

    # SHEET 7: Disruption scenarios + unmet demand
    ws6 = wb.create_sheet("Disruption & Unmet Demand")
    ws6.freeze_panes = "A2"

    dis_cols = ["Demand node", "Required demand"]
    for c in cases:
        dis_cols += [f"Unmet {c['label']}", f"Unmet % {c['label']}"]

    for ri, (c, cf) in enumerate(zip(cases, case_fills), 1):
        dis_facs = sorted(j for j in inst.V0 if c["z_worst"].get(j, 0) == 1)
        _set(ws6, ri, 1,
             f"Disrupted ({c['label']}): {dis_facs}",
             font=BOLD, fill=cf, align=left)
        ws6.row_dimensions[ri].height = 16

    start = len(cases) + 1
    _hrow(ws6, start, dis_cols,
          fills=[HDR_FILL, HDR_FILL] + [f for cf in case_fills for f in [cf, cf]])

    for ri, d in enumerate(sorted(inst.VD), start + 1):
        required = round(-inst.b[d], 2)
        bg = ALT_FILL if ri % 2 == 0 else None
        row_vals = [d, required]
        row_fmts = [None, NUM2]
        for c in cases:
            unmet = round(c["unmet_vals"].get(d, 0.0), 2)
            pct   = round(unmet / required * 100 if required > 0 else 0, 1)
            row_vals += [unmet, pct]
            row_fmts += [NUM2, "0.0"]
        for ci, (v, fmt) in enumerate(zip(row_vals, row_fmts), 1):
            fill = bg
            if ci >= 3 and (ci - 3) % 2 == 0:
                unmet_v = row_vals[ci - 1]
                fill = BAD_FILL if (isinstance(unmet_v, float) and unmet_v > 1e-4) else GOOD_FILL
            _set(ws6, ri, ci, v, fill=fill, align=center, num_fmt=fmt)
        ws6.row_dimensions[ri].height = 16

    _auto_col(ws6)

    # SHEET 8: Arc flow to demand nodes
    ws7 = wb.create_sheet("Flow to Demand (Arcs)")
    ws7.freeze_panes = "A2"

    flow_cols = ["Origin", "Dest", "Origin type",
                 f"Flow {c1['label']}", f"Cost {c1['label']}",
                 f"Flow {c2['label']}", f"Cost {c2['label']}",
                 f"Flow {c3['label']}", f"Cost {c3['label']}",
                 f"Flow {c4['label']}", f"Cost {c4['label']}"]
    _hrow(ws7, 1, flow_cols,
          fills=[HDR_FILL] * 3 + [C1_FILL] * 2 + [C2_FILL] * 2
                + [C3_FILL] * 2 + [C4_FILL] * 2)

    all_arcs = set()
    for c in cases:
        for row in c["flow_to_demand"]:
            all_arcs.add((row["origin"], row["dest"]))

    def _build_lookup(ftd):
        return {(r["origin"], r["dest"]): r for r in ftd}

    lk1 = _build_lookup(c1["flow_to_demand"])
    lk2 = _build_lookup(c2["flow_to_demand"])
    lk3 = _build_lookup(c3["flow_to_demand"])
    lk4 = _build_lookup(c4["flow_to_demand"])

    for ri, (orig, dest) in enumerate(sorted(all_arcs), 2):
        bg = ALT_FILL if ri % 2 == 0 else None
        otype = "supply" if orig in set(inst.VS) else "transship"
        r1 = lk1.get((orig, dest), {})
        r2 = lk2.get((orig, dest), {})
        r3 = lk3.get((orig, dest), {})
        r4 = lk4.get((orig, dest), {})
        vals = [orig, dest, otype,
                r1.get("flow"), r1.get("arc_cost"),
                r2.get("flow"), r2.get("arc_cost"),
                r3.get("flow"), r3.get("arc_cost"),
                r4.get("flow"), r4.get("arc_cost")]
        fmts = [None, None, None,
                NUM4, NUM2, NUM4, NUM2, NUM4, NUM2, NUM4, NUM2]
        for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
            cf = (C1_FILL if ci in (4, 5) else
                  C2_FILL if ci in (6, 7) else
                  C3_FILL if ci in (8, 9) else
                  C4_FILL if ci in (10, 11) else bg)
            _set(ws7, ri, ci, v,
                 fill=cf if v is not None else bg,
                 align=center, num_fmt=fmt)
        ws7.row_dimensions[ri].height = 16

    _auto_col(ws7)

    # SHEET 9: Origin node flow summary
    ws8 = wb.create_sheet("Origin Flow Summary")
    ws8.freeze_panes = "A2"

    sum_cols = ["Node", "Type",
                f"Flow {c1['label']}", f"Cost {c1['label']}",
                f"Flow {c2['label']}", f"Cost {c2['label']}",
                f"Flow {c3['label']}", f"Cost {c3['label']}",
                f"Flow {c4['label']}", f"Cost {c4['label']}",
                "Fixed cost", "Capacity"]
    _hrow(ws8, 1, sum_cols,
          fills=[HDR_FILL] * 2
                + [C1_FILL] * 2 + [C2_FILL] * 2
                + [C3_FILL] * 2 + [C4_FILL] * 2
                + [HDR_FILL] * 2)

    all_origins = sorted(set(
        r["node"] for c in cases for r in c["origin_summary"]))

    for ri, node in enumerate(all_origins, 2):
        ntype = "supply" if node in set(inst.VS) else "transship"
        type_fill = C1_FILL if ntype == "supply" else C2_FILL
        row_fill  = ALT_FILL if ri % 2 == 0 else None

        ol1 = {r["node"]: r for r in c1["origin_summary"]}
        ol2 = {r["node"]: r for r in c2["origin_summary"]}
        ol3 = {r["node"]: r for r in c3["origin_summary"]}
        ol4 = {r["node"]: r for r in c4["origin_summary"]}
        r1 = ol1.get(node, {})
        r2 = ol2.get(node, {})
        r3 = ol3.get(node, {})
        r4 = ol4.get(node, {})

        vals = [
            node, ntype,
            r1.get("total_flow"), r1.get("total_cost"),
            r2.get("total_flow"), r2.get("total_cost"),
            r3.get("total_flow"), r3.get("total_cost"),
            r4.get("total_flow"), r4.get("total_cost"),
            round(inst.f.get(node, 0), 2),
            round(inst.Q.get(node, 0), 2),
        ]
        fmts = [None, None,
                NUM4, NUM2, NUM4, NUM2, NUM4, NUM2, NUM4, NUM2,
                NUM2, NUM2]
        for ci, (v, fmt) in enumerate(zip(vals, fmts), 1):
            fill = type_fill if ci in (1, 2) else row_fill
            _set(ws8, ri, ci, v, fill=fill, align=center, num_fmt=fmt)
        ws8.row_dimensions[ri].height = 16

    tr = len(all_origins) + 2
    _set(ws8, tr, 1, "TOTAL", font=BOLD, align=left)
    for ci, c in zip([4, 6, 8, 10], cases):
        tot = round(sum(r["total_cost"] for r in c["origin_summary"]), 2)
        _set(ws8, tr, ci, tot, font=BOLD,
             fill=case_fills[[4, 6, 8, 10].index(ci)],
             align=center, num_fmt=NUM2)

    _auto_col(ws8)

    wb.save(out_path)
    print(f"\n  Excel comparison saved -> {out_path}")


# MAIN

def main():
    print("=" * 72)
    print("  FOUR-METHOD COMPARISON")
    print("  Case 1: Exact C&CG")
    print("  Case 2: Early-stopped C&CG (Phase 1 of H1)")
    print("  Case 3: H1 (C&CG + Phase 2)")
    print("  Case 4: H2 (MIP + Phase 2)")
    print("=" * 72)

    GEN_PARAMS = dict(
        nS       = 10,
        nT       = 20,
        nD       = 30,
        density  = 0.20,
        seed     = 4,
        m_budget = 5,
    )

    STOP_GAP_PCT     = 10.0
    EPS_EXACT        = 1e-6
    TIME_LIMIT_EXACT = 10800.0
    TIME_LIMIT_H1    = 10800.0
    TIME_LIMIT_H2    = 10800.0
    POOL_CAP_H1      = 0
    POOL_CAP_H2      = 0

    print(f"\nInstance params: {GEN_PARAMS}")
    print(f"Stop gap %     : {STOP_GAP_PCT}")
    print(f"Time limits    : exact={TIME_LIMIT_EXACT}s  H1={TIME_LIMIT_H1}s  H2={TIME_LIMIT_H2}s")
    print(f"Pool cap       : H1={POOL_CAP_H1 if POOL_CAP_H1 > 0 else 'unlimited'}  "
          f"H2={POOL_CAP_H2 if POOL_CAP_H2 > 0 else 'unlimited'}")
    raw = generate_instance_paper_exact(**GEN_PARAMS)
    summarize_instance(raw)
    inst = LNDPInstance.from_generator_dict(raw)

    print(f"\nNetwork: {len(inst.V)} nodes | {len(inst.arcs)} arcs | "
          f"{len(inst.V0)} facility nodes | m={inst.m}")

    results = run_comparison(
        inst,
        raw,
        gen_params       = GEN_PARAMS,
        stop_gap_pct     = STOP_GAP_PCT,
        time_limit_exact = TIME_LIMIT_EXACT,
        time_limit_H1    = TIME_LIMIT_H1,
        time_limit_H2    = TIME_LIMIT_H2,
        eps_exact        = EPS_EXACT,
        pool_cap_H1      = POOL_CAP_H1,
        pool_cap_H2      = POOL_CAP_H2,
        verbose          = True,
    )

    print_comparison(results, inst)

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    label = (f"{int(GEN_PARAMS['density']*100)}pct-"
             f"{GEN_PARAMS['nS']}-{GEN_PARAMS['nT']}-{GEN_PARAMS['nD']}")
    out_path = (f"COMPARE_{label}_seed{GEN_PARAMS['seed']}"
                f"_m{GEN_PARAMS['m_budget']}_{ts}.xlsx")

    export_comparison_excel(results, inst, raw, GEN_PARAMS, out_path)


if __name__ == "__main__":
    main()
